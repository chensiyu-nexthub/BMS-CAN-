from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path

p = Path(r'D:\xianmu\BMS-main\开题报告初版_修改版_已修正.docx')
d = Document(p)
target = None
for para in d.paragraphs:
    if para.text.strip().startswith('图 1 系统框图') or para.text.strip().startswith('图1系统框图'):
        target = para
        break

print('caption:', target.text if target else None)
if target:
    new = d.add_paragraph()
    target._p.addprevious(new._p)
    new.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = new.add_run()
    run.add_picture(r'D:\xianmu\BMS-main\开题提示词渲染检查\系统框图原始图_rgb.png', width=Inches(4.5))
    d.save(p)
    print('saved:', p)
