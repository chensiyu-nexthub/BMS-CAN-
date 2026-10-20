from docx import Document
import re

p = r"D:\xianmu\BMS-main\桌面补充图片检查_格式修正版.docx"
d = Document(p)

issues = []
for i, x in enumerate(d.paragraphs):
    t = x.text
    if not t.strip():
        continue
    pattern = r"[\u4e00-\u9fff]\s+[A-Za-z0-9]|[A-Za-z0-9]\s+[\u4e00-\u9fff]"
    for m in re.finditer(pattern, t):
        issues.append((i, "CJK_ASCII_SPACE", m.group(0), t[max(0, m.start() - 18):m.end() + 18]))
    for pat in ["， ", "。 ", "； ", "： ", " ，", " 。"]:
        start = 0
        while True:
            pos = t.find(pat, start)
            if pos < 0:
                break
            issues.append((i, "PUNCT_SPACE", pat, t[max(0, pos - 18):pos + len(pat) + 18]))
            start = pos + 1

print("space_issue_count", len(issues))
for it in issues[:120]:
    print(it)

print("ASCII_PUNCT_BODY")
for i, x in enumerate(d.paragraphs):
    t = x.text
    if 33 < i < 201 and t.strip():
        for ch in ",;:!?":
            if ch in t:
                pos = t.find(ch)
                print(i, repr(ch), t[max(0, pos - 30):pos + 30])
                break

print("QUOTE_PARAS")
for i, x in enumerate(d.paragraphs):
    if any(q in x.text for q in ["“", "”", "‘", "’", "\""]):
        print(i, x.text[:150])

print("HEADINGS_AND_CAPTIONS")
for i, x in enumerate(d.paragraphs):
    t = x.text.strip()
    if not t:
        continue
    if (re.match(r"^[一二三四五六七八九十]+、", t) or t.startswith("第") or
            re.match(r"^图\d+", t) or t.startswith("研究") or t.endswith("：") or t.endswith(":")):
        print(i, repr(t))

print("EMPTY_PARA_AROUND_SECTIONS")
for i, x in enumerate(d.paragraphs):
    t = x.text.strip()
    if t and (re.match(r"^[一二三四五六七八九十]+、", t) or t.startswith("图")):
        print(i, t[:60])

print("PARAGRAPH_COUNT", len(d.paragraphs))
print("SECTION_TEXT_COUNT")
for i, x in enumerate(d.paragraphs):
    t = x.text.strip()
    if t:
        print(i, len(t), t[:70])
