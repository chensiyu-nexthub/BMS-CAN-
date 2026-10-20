# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

SOURCE = r"D:\xianmu\baogao\本科毕业论文开题报告_最终合并版.docx"
OUTPUT = r"D:/xianmu/BMS-main/本科毕业论文开题报告_提示词模板_学校五项严格原样重排版.docx"

src = Document(SOURCE)
texts = [p.text for p in src.paragraphs if p.text.strip()]

old_level1 = {
    "一、研究目的和意义",
    "二、国内外研究现状分析",
    "三、研究（设计）主要内容、预期目标及拟解决的关键问题",
    "四、研究方法与技术路线",
    "五、主要参考文献",
    "六、进度安排",
    "七、大纲",
}

# 原版非空段落索引：0标题，1总约束；2-84第一部分；85-94第二部分；
# 95-101第三部分；102-103参考文献；104-105进度；106-107大纲；108结尾。
groups = [
    ("一、综述本课题国内外研究动态，说明选题的依据和意义", list(range(2, 85))),
    ("二、研究的基本内容，拟解决的主要问题", list(range(85, 95))),
    ("三、研究步骤、方法及措施", list(range(95, 102))),
    ("四、研究工作进度", [104, 105]),
    ("五、主要参考文献", [102, 103]),
    ("附加、大纲（不属于学校模板五项，需要时单独使用）", [106, 107]),
]

new = Document()
for section in new.sections:
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2.5)

style = new.styles["Normal"]
style.font.name = "宋体"
style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
style.font.size = Pt(12)

def set_font(run, size=12, bold=False):
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold

def add_para(text, size=12, bold=False):
    p = new.add_paragraph()
    r = p.add_run(text)
    set_font(r, size, bold)
    p.paragraph_format.line_spacing = Pt(20)
    return p

p = new.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("本科毕业论文（设计）开题报告AI撰写提示词（学校模板严格原样重排版）")
set_font(r, 16, True)

add_para(texts[1])

for school_title, indices in groups:
    add_para(school_title, 14, True)
    for i in indices:
        if texts[i] in old_level1:
            continue
        add_para(texts[i])

add_para(texts[108])

new.save(OUTPUT)
print(OUTPUT)
