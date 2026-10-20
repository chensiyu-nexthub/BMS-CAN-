from docx import Document
from docx.shared import Pt
import re

p = r"D:\xianmu\BMS-main\桌面补充图片检查_格式修正版.docx"
d = Document(p)

print("STYLE_AND_FORMAT")
for i, x in enumerate(d.paragraphs):
    t = x.text.strip()
    if not t:
        continue
    if 32 < i < 161 and not t.startswith("["):
        style = x.style.name
        line_rule = None
        line_spacing = None
        if x.paragraph_format.line_spacing_rule is not None:
            line_rule = str(x.paragraph_format.line_spacing_rule)
        if x.paragraph_format.line_spacing is not None:
            line_spacing = x.paragraph_format.line_spacing
        fonts = []
        sizes = []
        for r in x.runs:
            if r.text:
                fonts.append(r.font.name or (r.style.font.name if r.style and r.style.font else None))
                if r.font.size:
                    sizes.append(r.font.size.pt)
        fonts = sorted(set(z for z in fonts if z))
        sizes = sorted(set(sizes))
        if fonts != ["宋体"] or sizes != [12.0] or line_rule != "EXACTLY" or line_spacing != 20:
            print(i, t[:42], "style=", style, "fonts=", fonts, "sizes=", sizes,
                  "line=", line_rule, line_spacing, "first_indent=", x.paragraph_format.first_line_indent,
                  "alignment=", x.paragraph_format.alignment)

print("SECTION_PAGE")
for i, s in enumerate(d.sections):
    print(i, s.page_width, s.page_height, s.left_margin, s.right_margin, s.top_margin, s.bottom_margin)

print("TABLE_PARAGRAPH_COUNT", len(d.tables))
print("BODY_CHARS", sum(len(x.text.strip()) for x in d.paragraphs[32:161]))
